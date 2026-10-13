"""Punto de entrada del proceso local de licitaciones adjudicadas.

Coordina las etapas existentes, sin duplicar su logica:
    1. (solo con --actualizar) Ingesta/extract.py con las credenciales de Ingesta/.env
    2. Limpieza y Excel intermedio (generar_excel_intermedio.generar)
    3. Verificacion del Excel generado y de que los CSV de origen no cambiaron

Modos:
    (por defecto)   DEMOSTRACION: usa solo los CSV ya guardados en Ingesta/output.
                    No consulta Licitaciones.info ni necesita credenciales.
    --actualizar    ACTUALIZACION: ejecuta primero la ingesta. Si falla, o no deja
                    exactamente un CSV nuevo, el proceso se detiene y NO genera Excel.

Opciones:
    --salida DIR    carpeta del Excel (por defecto datos/intermedio)

Codigos de salida: 0 correcto | 1 faltan CSV o configuracion | 2 fallo la ingesta |
3 la verificacion del Excel o de los CSV fallo (el Excel se renombra con sufijo _REVISAR).

Uso (desde la raiz del proyecto):
    .venv/bin/python Transformacion/ejecutar_proceso.py
    .venv/bin/python Transformacion/ejecutar_proceso.py --actualizar
"""
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
from pathlib import Path

from dotenv import dotenv_values
from openpyxl import load_workbook

import generar_excel_intermedio as gen
import limpieza_licitaciones as lim

RAIZ = gen.RAIZ
INGESTA = RAIZ / "Ingesta"
ENTRADA = gen.ENTRADA_POR_DEFECTO
ARCHIVO_ENV = INGESTA / ".env"

# Columnas que deben quedar vacias mientras no haya una fuente confirmada.
SIN_FUENTE_DATOS = [c.nombre for c in lim.COLUMNAS_GRECO if c.estado == lim.SIN_FUENTE]
SIN_FUENTE_GANADORES = [c.nombre for c in lim.COLUMNAS_GANADORES if c.estado == lim.SIN_FUENTE]
PATRON_JWT = re.compile(r"eyJ[A-Za-z0-9_-]{5,}\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+")
PATRON_CORREO = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PATRON_BEARER = re.compile(r"Bearer\s+\S+", re.IGNORECASE)


def aviso(texto: str = "") -> None:
    print(texto, flush=True)


# ----------------------------------------------------------------------
# Configuracion de la ingesta (sin mostrar valores secretos)
# ----------------------------------------------------------------------
def valor_efectivo(nombre: str, archivo_env: dict) -> str:
    """Igual que config.py: una variable ya definida en el entorno prevalece sobre el .env."""
    if nombre in os.environ:
        return os.environ[nombre] or ""
    return archivo_env.get(nombre) or ""


def secretos_locales(archivo_env: dict) -> list[str]:
    return [v for v in (valor_efectivo(k, archivo_env) for k in ("LICITACIONES_EMAIL", "LICITACIONES_PASSWORD")) if v]


def ocultar(linea: str, secretos: list[str]) -> str:
    for secreto in secretos:
        linea = linea.replace(secreto, "<oculto>")
    linea = PATRON_JWT.sub("<JWT oculto>", linea)
    linea = PATRON_BEARER.sub("Bearer <oculto>", linea)
    return PATRON_CORREO.sub("<correo oculto>", linea)


def comprobar_configuracion() -> dict | None:
    if not ARCHIVO_ENV.exists():
        aviso(f"ERROR: no existe {ARCHIVO_ENV.relative_to(RAIZ)}. Copia Ingesta/.env.example y completa las credenciales.")
        return None
    archivo_env = dotenv_values(ARCHIVO_ENV)
    faltan = [k for k in ("LICITACIONES_EMAIL", "LICITACIONES_PASSWORD") if not valor_efectivo(k, archivo_env)]
    if faltan:
        aviso(f"ERROR: faltan valores en Ingesta/.env: {', '.join(faltan)}.")
        return None
    estado = valor_efectivo("LICITACIONES_ESTADO", archivo_env)
    if estado != lim.ESTADO_ADJUDICADO:
        aviso(
            f"ERROR: LICITACIONES_ESTADO es '{estado or '(vacio)'}' y este proceso esta limitado a "
            f"'{lim.ESTADO_ADJUDICADO}'. Ajusta Ingesta/.env antes de actualizar."
        )
        return None
    aviso(f"Configuracion: credenciales definidas (valores ocultos); filtro de la ingesta LICITACIONES_ESTADO={estado}.")
    return archivo_env


# ----------------------------------------------------------------------
# Etapas
# ----------------------------------------------------------------------
def ejecutar_ingesta(secretos: list[str]) -> Path | None:
    """Ejecuta Ingesta/extract.py tal como esta. Devuelve el CSV nuevo o None si fallo."""
    antes = {p.name for p in ENTRADA.glob("licitaciones_*.csv")}
    aviso("Ejecutando Ingesta/extract.py (salida de la ingesta, con datos sensibles ocultos):")
    with subprocess.Popen(
        [sys.executable, "extract.py"],
        cwd=INGESTA,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
    ) as proceso:
        assert proceso.stdout is not None
        for linea in proceso.stdout:
            aviso("  | " + ocultar(linea.rstrip("\n"), secretos))
        codigo = proceso.wait()

    if codigo != 0:
        aviso(
            f"\nERROR: la ingesta termino con codigo {codigo}. El proceso se detiene y NO se genera Excel.\n"
            "Revisa el mensaje anterior: credenciales rechazadas (401), bloqueo de Cloudflare/WAF, "
            "falta de conexion o un error de la API."
        )
        return None

    nuevos = sorted({p.name for p in ENTRADA.glob("licitaciones_*.csv")} - antes)
    if len(nuevos) != 1:
        aviso(
            f"\nERROR: la ingesta termino sin error pero dejo {len(nuevos)} CSV nuevos (se esperaba 1). "
            "El proceso se detiene y NO se genera Excel."
        )
        return None
    return ENTRADA / nuevos[0]


def verificar(destino: Path, huellas_antes: dict[str, str]) -> list[str]:
    """Comprueba el Excel generado y que los CSV de origen siguen intactos."""
    problemas: list[str] = []

    for archivo in lim.listar_csvs(ENTRADA):
        anterior = huellas_antes.get(archivo.ruta.name)
        if anterior and anterior != archivo.sha256:
            problemas.append(f"El CSV {archivo.ruta.name} cambio durante el proceso.")

    libro = load_workbook(destino, read_only=True)
    for hoja, vacias in (("Datos", SIN_FUENTE_DATOS), ("Ganadores", SIN_FUENTE_GANADORES)):
        filas = libro[hoja].iter_rows(values_only=True)
        encabezados = list(next(filas))
        registros = [dict(zip(encabezados, fila)) for fila in filas]

        for columna in vacias:
            llenas = sum(1 for r in registros if r.get(columna) not in (None, ""))
            if llenas:
                problemas.append(f"{hoja}: la columna sin fuente '{columna}' tiene {llenas} valores.")
        for r in registros:
            if any(isinstance(v, str) and PATRON_JWT.search(v) for v in r.values()):
                problemas.append(f"{hoja}: hay un valor con forma de JWT.")
                break

        if hoja == "Datos":
            estados = {r.get("Estado (origen)") for r in registros}
            ids = [r.get("idContrato") for r in registros]
            if estados - {lim.ESTADO_ADJUDICADO}:
                problemas.append(f"Datos: hay estados distintos de '{lim.ESTADO_ADJUDICADO}': {sorted(map(str, estados))}.")
            if len(ids) != len(set(ids)):
                problemas.append("Datos: hay idContrato duplicados.")
            total_datos = len(registros)
    libro.close()

    if not problemas:
        aviso(
            f"Verificacion: OK. {total_datos} procesos en Datos, todos con estado '{lim.ESTADO_ADJUDICADO}', "
            "sin idContrato duplicados; NIT, Localidad, Latitud, Longitud, Contrato No, Fecha Adjudicacion, "
            "Nombre de la Obra e Integrantes del consorcio vacios; sin JWT; CSV de origen sin cambios."
        )
    return problemas


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument(
        "--actualizar",
        action="store_true",
        help="ejecuta primero la ingesta desde Licitaciones.info; si falla, no genera Excel",
    )
    parser.add_argument("--salida", type=Path, default=gen.SALIDA_POR_DEFECTO, help="carpeta del Excel (por defecto datos/intermedio)")
    args = parser.parse_args()

    if args.actualizar:
        aviso("=== MODO ACTUALIZACION: se consulta Licitaciones.info con las credenciales de Ingesta/.env ===")
        archivo_env = comprobar_configuracion()
        if archivo_env is None:
            return 1
        nuevo = ejecutar_ingesta(secretos_locales(archivo_env))
        if nuevo is None:
            return 2
        fecha = lim.fecha_desde_nombre(nuevo)
        filas = len(lim.leer_filas(lim.ArchivoFuente(nuevo, fecha, "")))
        aviso(f"\nIngesta correcta: {nuevo.relative_to(RAIZ)} ({filas} filas, extraido el {fecha:%d/%m/%Y %H:%M:%S}).")
        if filas == 0:
            aviso("La consulta no devolvio procesos; el Excel consolidara los CSV que ya estaban guardados.")
        modo = (
            f"Actualización: ingesta ejecutada y guardada en {nuevo.name} ({filas} filas, extraída el "
            f"{fecha:%d/%m/%Y %H:%M:%S}); consolidada con los demás CSV locales."
        )
    else:
        aviso("=== MODO DEMOSTRACION: se usan CSV ya guardados localmente; NO se consulta Licitaciones.info ===")
        modo = "Demostración: se usaron CSV ya guardados localmente; no se consultó Licitaciones.info."

    archivos = lim.listar_csvs(ENTRADA)
    if not archivos:
        aviso(f"ERROR: no hay CSV en {ENTRADA.relative_to(RAIZ)}. Ejecuta con --actualizar para obtenerlos.")
        return 1
    aviso(f"\nCSV que se consolidan ({len(archivos)}), del mas antiguo al mas reciente:")
    for archivo in archivos:
        fecha = f"{archivo.fecha_extraccion:%d/%m/%Y %H:%M:%S}" if archivo.fecha_extraccion else "fecha no legible"
        aviso(f"  - {archivo.ruta.relative_to(RAIZ)} (extraido el {fecha})")
    huellas_antes = {a.ruta.name: a.sha256 for a in archivos}

    try:
        destino, resultado = gen.generar(ENTRADA, args.salida, modo)
    except FileExistsError as exc:
        aviso(f"ERROR: {exc.filename} ya existe; no se sobrescribe. Vuelve a ejecutar en unos segundos.")
        return 1

    m = resultado.metricas
    aviso(
        f"\nLimpieza: {m['filas_leidas']} filas leidas, {m['procesos_unicos']} procesos unicos, "
        f"{m['incluidos']} adjudicados incluidos, {m['excluidos_por_estado']} excluidos por estado, "
        f"{m['filas_ganadores']} ganadores, {len(resultado.incidencias)} incidencias."
    )

    problemas = verificar(destino, huellas_antes)
    if problemas:
        revisar = destino.with_name(destino.stem + "_REVISAR.xlsx")
        destino.rename(revisar)
        aviso("Verificacion: FALLO. No uses este Excel sin revisarlo:")
        for problema in problemas:
            aviso(f"  - {problema}")
        aviso(f"Excel marcado para revision: {revisar.relative_to(RAIZ) if revisar.is_relative_to(RAIZ) else revisar}")
        return 3

    aviso(f"\nExcel generado: {destino.relative_to(RAIZ) if destino.is_relative_to(RAIZ) else destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
