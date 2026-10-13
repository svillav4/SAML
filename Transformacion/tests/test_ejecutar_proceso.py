"""Pruebas del punto de entrada ejecutar_proceso.py, SIN consultar Licitaciones.info.

La ingesta real se sustituye por un extract.py falso dentro de una carpeta
temporal; las credenciales y el JWT son valores sinteticos.

Uso (desde la raiz del proyecto):
    .venv/bin/python -m unittest discover -s Transformacion/tests -v
"""
from __future__ import annotations

import contextlib
import io
import os
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

CARPETA = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(CARPETA))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import ejecutar_proceso as ep  # noqa: E402
from test_limpieza import escribir_csv, proceso  # noqa: E402

CORREO_FALSO = "persona.prueba@ejemplo.test"
CLAVE_FALSA = "Clave-Sintetica-123"
JWT_FALSO = "eyJhbGciOiJIUzI1NiJ9.eyJzdWIiOiJwcnVlYmEifQ.firmaSintetica"


class EjecutarProcesoTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        raiz = Path(self._tmp.name)
        self.ingesta = raiz / "Ingesta"
        self.output = self.ingesta / "output"
        self.output.mkdir(parents=True)
        self.salida = raiz / "datos" / "intermedio"
        self.env = self.ingesta / ".env"
        self.env.write_text(
            f"LICITACIONES_EMAIL={CORREO_FALSO}\nLICITACIONES_PASSWORD={CLAVE_FALSA}\nLICITACIONES_ESTADO=Adjudicado\n",
            encoding="utf-8",
        )
        escribir_csv(self.ingesta, "plantilla.csv", [proceso("1"), proceso("2", estado="Abierto")])

        variables = {k: v for k, v in os.environ.items() if not k.startswith("LICITACIONES_")}
        self._parches = [
            mock.patch.object(ep, "RAIZ", raiz),
            mock.patch.object(ep, "INGESTA", self.ingesta),
            mock.patch.object(ep, "ENTRADA", self.output),
            mock.patch.object(ep, "ARCHIVO_ENV", self.env),
            mock.patch.dict(os.environ, variables, clear=True),
        ]
        for parche in self._parches:
            parche.start()

    def tearDown(self) -> None:
        for parche in reversed(self._parches):
            parche.stop()
        self._tmp.cleanup()

    def _extract_falso(self, cuerpo: str) -> None:
        (self.ingesta / "extract.py").write_text(textwrap.dedent(cuerpo), encoding="utf-8")

    def _ejecutar(self, *argumentos: str) -> tuple[int, str]:
        salida = io.StringIO()
        with contextlib.redirect_stdout(salida), mock.patch.object(
            sys, "argv", ["ejecutar_proceso.py", "--salida", str(self.salida), *argumentos]
        ):
            codigo = ep.main()
        return codigo, salida.getvalue()

    def _excels(self) -> list[Path]:
        return sorted(self.salida.glob("*.xlsx")) if self.salida.exists() else []

    def test_si_la_ingesta_falla_no_genera_excel_ni_muestra_secretos(self) -> None:
        escribir_csv(self.output, "licitaciones_20260101_080000.csv", [proceso("9")])
        self._extract_falso(
            f"""
            import sys
            print("Login con {CORREO_FALSO} / {CLAVE_FALSA}")
            print("Authorization: Bearer {JWT_FALSO}")
            print("token {JWT_FALSO}")
            sys.exit(1)
            """
        )

        codigo, texto = self._ejecutar("--actualizar")

        self.assertEqual(codigo, 2)
        self.assertEqual(self._excels(), [])
        self.assertIn("NO se genera Excel", texto)
        for secreto in (CORREO_FALSO, CLAVE_FALSA, JWT_FALSO):
            self.assertNotIn(secreto, texto)

    def test_si_la_ingesta_termina_bien_genera_excel_solo_con_adjudicados(self) -> None:
        self._extract_falso(
            """
            import shutil
            shutil.copy("plantilla.csv", "output/licitaciones_20260301_090000.csv")
            print("Se encontraron 2 licitaciones.")
            """
        )

        codigo, texto = self._ejecutar("--actualizar")

        self.assertEqual(codigo, 0, texto)
        self.assertIn("Ingesta correcta", texto)
        self.assertIn("Verificacion: OK", texto)
        (excel,) = self._excels()
        libro = load_workbook(excel)
        encabezados = [c.value for c in libro["Datos"][1]]
        columna_estado = encabezados.index("Estado (origen)")
        valores = [fila[columna_estado] for fila in libro["Datos"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(valores, ["Adjudicado"])
        modo = [f[1] for f in libro["Resumen"].iter_rows(values_only=True) if f[0] == "Modo de ejecución"]
        self.assertTrue(modo and modo[0].startswith("Actualización"))

    def test_actualizar_se_detiene_si_el_filtro_no_es_adjudicado(self) -> None:
        self.env.write_text(
            f"LICITACIONES_EMAIL={CORREO_FALSO}\nLICITACIONES_PASSWORD={CLAVE_FALSA}\nLICITACIONES_ESTADO=Abierto\n",
            encoding="utf-8",
        )
        self._extract_falso("raise SystemExit('no deberia ejecutarse')")

        codigo, texto = self._ejecutar("--actualizar")

        self.assertEqual(codigo, 1)
        self.assertIn("limitado a 'Adjudicado'", texto)
        self.assertNotIn(CLAVE_FALSA, texto)
        self.assertEqual(self._excels(), [])

    def test_demo_usa_csv_locales_y_no_ejecuta_la_ingesta(self) -> None:
        escribir_csv(self.output, "licitaciones_20260101_080000.csv", [proceso("1")])
        self._extract_falso("raise SystemExit('la demo no debe ejecutar la ingesta')")

        codigo, texto = self._ejecutar()

        self.assertEqual(codigo, 0, texto)
        self.assertIn("MODO DEMOSTRACION", texto)
        self.assertIn("licitaciones_20260101_080000.csv (extraido el 01/01/2026 08:00:00)", texto)
        self.assertEqual(len(self._excels()), 1)

    def test_demo_sin_csv_termina_con_error_y_sin_excel(self) -> None:
        codigo, texto = self._ejecutar()

        self.assertEqual(codigo, 1)
        self.assertIn("no hay CSV", texto)
        self.assertEqual(self._excels(), [])


if __name__ == "__main__":
    unittest.main()
