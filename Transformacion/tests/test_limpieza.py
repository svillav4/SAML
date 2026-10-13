"""Pruebas de la limpieza intermedia con datos SINTETICOS (no son datos reales).

Los CSV se escriben con el mismo formato que Ingesta/extract.py: listas
anidadas guardadas como texto de Python.

Uso (desde la raiz del proyecto):
    .venv/bin/python -m unittest discover -s Transformacion/tests -v
"""
from __future__ import annotations

import csv
import hashlib
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import load_workbook

CARPETA = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(CARPETA))

import limpieza_licitaciones as lim  # noqa: E402

JWT_FALSO = "eyJhbGciOiJIUzI1NiJ9.SINTETICO.FIRMA"


def ganador(id_contrato: str, id_proponente: int, nombre: str, consorcio=False) -> dict:
    return {
        "idContrato": id_contrato,
        "idProponente": id_proponente,
        "es_consorcio": consorcio,
        "nombreEntidad": nombre,
        "id_contratista_verificado": None,
        "jwt_contratista_id": JWT_FALSO,
        "pivot": {"idContrato": id_contrato, "idProponente": id_proponente},
    }


def proceso(id_contrato: str, estado: str = "Adjudicado", **cambios) -> dict:
    base = {
        "link": f"https://col.licitaciones.info/detalle/{id_contrato}",
        "nombre_fuente": "secop II",
        "idContrato": id_contrato,
        "Nombre": f"Licitacion publica PRUEBA-{id_contrato}",
        "CodigoProceso": f"PRUEBA-{id_contrato}",
        "FechaPublicacion": "2026-02-12 10:30:00",
        "EntidadContratante": "  ENTIDAD   DE PRUEBA ",
        "Objeto": "Obra de prueba",
        "FechaVencimiento": "2026-03-01 17:00:00",
        "Valor": "1500000000",
        "TextoDepartamento": "Cundinamarca : Bogotá D.C.",
        "estado_agrupado": estado,
        "leido": 0,
        "ganadores": [ganador(id_contrato, 1, "CONSORCIO PRUEBA", True)],
        "proponentes": [ganador(id_contrato, 1, "CONSORCIO PRUEBA", True)],
        "cronogramas": [{"nombre": "Audiencia de Adjudicación", "fecha": "20/03/2026 - 10:00 AM"}],
    }
    base.update(cambios)
    return base


def escribir_csv(carpeta: Path, nombre: str, procesos: list[dict]) -> Path:
    columnas = list(procesos[0].keys())
    ruta = carpeta / nombre
    with ruta.open("w", newline="", encoding="utf-8") as f:
        escritor = csv.DictWriter(f, fieldnames=columnas)
        escritor.writeheader()
        for p in procesos:
            escritor.writerow({k: p.get(k, "") for k in columnas})
    return ruta


class LimpiezaTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_consolida_por_idcontrato_con_la_version_mas_reciente(self) -> None:
        escribir_csv(self.dir, "licitaciones_20260101_080000.csv", [proceso("1", Valor="100000000"), proceso("2")])
        escribir_csv(self.dir, "licitaciones_20260201_080000.csv", [proceso("1", Valor="200000000")])

        r = lim.procesar(self.dir)

        self.assertEqual(r.metricas["filas_leidas"], 3)
        self.assertEqual(r.metricas["reemplazados_por_version_reciente"], 1)
        self.assertEqual(sorted(d["idContrato"] for d in r.datos), ["1", "2"])
        uno = next(d for d in r.datos if d["idContrato"] == "1")
        self.assertEqual(uno["Valor Millones"], 200.0)
        self.assertEqual(uno["Archivo origen"], "licitaciones_20260201_080000.csv")
        self.assertTrue(any(i["Tipo"] == lim.INC_DUPLICADO for i in r.incidencias))

    def test_excluye_no_adjudicados_y_lo_registra(self) -> None:
        escribir_csv(self.dir, "licitaciones_20260101_080000.csv", [proceso("1"), proceso("2", estado="Abierto")])

        r = lim.procesar(self.dir)

        self.assertEqual([d["idContrato"] for d in r.datos], ["1"])
        self.assertEqual(r.metricas["excluidos_por_estado"], 1)
        excluidos = [i for i in r.incidencias if i["Tipo"] == lim.INC_EXCLUIDO]
        self.assertEqual([i["idContrato"] for i in excluidos], ["2"])

    def test_varios_ganadores_y_repetidos_sin_perder_informacion(self) -> None:
        ganadores = [ganador("1", 10, "EMPRESA A"), ganador("1", 11, "UNION TEMPORAL B", True), ganador("1", 10, "EMPRESA A")]
        escribir_csv(self.dir, "licitaciones_20260101_080000.csv", [proceso("1", ganadores=ganadores, proponentes=ganadores[:1])])

        r = lim.procesar(self.dir)

        self.assertEqual(r.datos[0]["Empresa"], "EMPRESA A | UNION TEMPORAL B")
        self.assertEqual([g["Empresa"] for g in r.ganadores], ["EMPRESA A", "UNION TEMPORAL B"])
        self.assertEqual([g["Repeticiones en origen"] for g in r.ganadores], [2, 1])
        self.assertEqual([g["Está en proponentes"] for g in r.ganadores], ["Sí", "No"])
        self.assertIsNone(r.ganadores[1]["Integrantes del consorcio"])

    def test_campos_vacios_no_se_inventan_ni_eliminan_el_registro(self) -> None:
        escribir_csv(
            self.dir,
            "licitaciones_20260101_080000.csv",
            [proceso("1", FechaVencimiento="", Valor="0", ganadores=[], cronogramas=[])],
        )

        r = lim.procesar(self.dir)

        self.assertEqual(len(r.datos), 1)
        fila = r.datos[0]
        self.assertIsNone(fila["Fecha vencimiento (origen)"])
        self.assertIsNone(fila["Valor Millones"])
        self.assertEqual(fila["Empresa"], "")
        self.assertIsNone(fila["Fecha Adjudicación"])
        campos = {i["Campo"] for i in r.incidencias if i["Alcance"] == "Registro"}
        self.assertTrue({"FechaVencimiento", "Valor", "ganadores"} <= campos)

    def test_fecha_adjudicacion_no_se_toma_del_cronograma(self) -> None:
        escribir_csv(self.dir, "licitaciones_20260101_080000.csv", [proceso("1")])

        fila = lim.procesar(self.dir).datos[0]

        self.assertIsNone(fila["Fecha Adjudicación"])
        self.assertEqual(fila["Adjudicación programada (cronograma)"], "20/03/2026 10:00")
        self.assertEqual(fila["Entidad"], "ENTIDAD DE PRUEBA")

    def test_excel_sin_campos_excluidos_y_sin_sobrescribir(self) -> None:
        entrada = self.dir / "output"
        entrada.mkdir()
        csv_origen = escribir_csv(entrada, "licitaciones_20260101_080000.csv", [proceso("1"), proceso("2")])
        hash_antes = hashlib.sha256(csv_origen.read_bytes()).hexdigest()
        salida = self.dir / "intermedio"
        comando = [
            sys.executable,
            str(CARPETA / "generar_excel_intermedio.py"),
            "--entrada",
            str(entrada),
            "--salida",
            str(salida),
        ]

        subprocess.run(comando, check=True, capture_output=True, text=True)

        self.assertEqual(hashlib.sha256(csv_origen.read_bytes()).hexdigest(), hash_antes)
        (libro_ruta,) = salida.glob("licitaciones_intermedio_*.xlsx")
        libro = load_workbook(libro_ruta)
        self.assertEqual(libro.sheetnames, ["Datos", "Ganadores", "Incidencias", "Resumen"])
        self.assertEqual(libro["Datos"].max_row, 3)
        textos = [str(c.value) for hoja in libro for fila in hoja.iter_rows() for c in fila if c.value is not None]
        self.assertFalse(any(JWT_FALSO in t for t in textos))
        encabezados = {c.value for c in libro["Datos"][1]} | {c.value for c in libro["Ganadores"][1]}
        self.assertFalse(encabezados & {"jwt_contratista_id", "pivot", "leido"})

        # Una nueva ejecucion crea otro archivo y no toca el anterior.
        contenido_antes = libro_ruta.read_bytes()
        time.sleep(1.1)
        subprocess.run(comando, check=True, capture_output=True, text=True)
        self.assertEqual(len(list(salida.glob("licitaciones_intermedio_*.xlsx"))), 2)
        self.assertEqual(libro_ruta.read_bytes(), contenido_antes)


if __name__ == "__main__":
    unittest.main()
