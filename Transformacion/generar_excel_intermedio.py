"""Genera el Excel intermedio de licitaciones adjudicadas a partir de la ingesta.

Lee todos los CSV licitaciones_*.csv de Ingesta/output (sin modificarlos), los
consolida por idContrato con limpieza_licitaciones.py y escribe un libro nuevo
con las hojas Datos, Ganadores, Incidencias y Resumen en
datos/intermedio/licitaciones_intermedio_<fecha_hora>.xlsx. Nunca sobrescribe
un archivo existente.

Es una version local de demostracion y validacion: no reemplaza la
integracion final con Obras V6, Supabase ni n8n.

Uso (desde la raiz del proyecto):
    .venv/bin/python Transformacion/generar_excel_intermedio.py
    .venv/bin/python Transformacion/generar_excel_intermedio.py --entrada Ingesta/output --salida datos/intermedio
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import limpieza_licitaciones as lim

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA_POR_DEFECTO = RAIZ / "Ingesta" / "output"
SALIDA_POR_DEFECTO = RAIZ / "datos" / "intermedio"

COLORES = {
    lim.MAPEADO: "C6EFCE",
    lim.PROVISIONAL: "FFEB9C",
    lim.SIN_FUENTE: "F8CBAD",
    lim.MANUAL: "D9D9D9",
    lim.CALCULADO: "E4DFEC",
    lim.TRAZABILIDAD: "DDEBF7",
}
SIGNIFICADO_COLORES = {
    lim.MAPEADO: "Se llena desde la API con una regla confirmada.",
    lim.PROVISIONAL: "Se llena desde la API, pero el mapeo debe confirmarse con Obras V6.",
    lim.SIN_FUENTE: "La API no trae el dato o no se identifica con certeza; queda vacía.",
    lim.MANUAL: "La diligencia el equipo comercial de GRECO; queda vacía.",
    lim.CALCULADO: "Depende de fórmulas o datos del archivo de GRECO; queda vacía.",
    lim.TRAZABILIDAD: "Datos de origen para revisar los mapeos y relacionar registros.",
}

ANCHOS = {
    "Entidad": 38,
    "Empresa": 45,
    "Nombre de la Obra": 30,
    "Despacho/Disposición": 16,
    "Enlace": 30,
    "Nombre (origen)": 40,
    "Objeto (origen)": 70,
    "Ubicación (origen)": 30,
    "Adjudicación programada (cronograma)": 30,
    "Archivo origen": 34,
    "Detalle": 90,
    "Campo": 24,
    "Tipo": 24,
}
FORMATOS = {
    "Valor Millones": "#,##0.00",
    "Valor (origen, pesos)": "#,##0",
    "Fecha publicación (origen)": "dd/mm/yyyy hh:mm",
    "Fecha vencimiento (origen)": "dd/mm/yyyy hh:mm",
    "Fecha extracción": "dd/mm/yyyy hh:mm",
}

FUENTE_ENCABEZADO = Font(bold=True, color="1F1F1F")
FUENTE_TITULO = Font(bold=True, size=14)
FUENTE_SECCION = Font(bold=True, size=12, color="1F4E78")
BORDE = Border(bottom=Side(style="thin", color="808080"))
ALINEACION_ENCABEZADO = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _celda(valor: Any) -> Any:
    """openpyxl rechaza caracteres de control en los textos."""
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    return valor


def _escribir_tabla(
    hoja: Worksheet,
    columnas: list[str],
    filas: Iterable[dict[str, Any]],
    *,
    estados: dict[str, lim.Columna] | None = None,
    congelar: str = "A2",
) -> None:
    hoja.append(columnas)
    for indice, nombre in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice)
        celda.font = FUENTE_ENCABEZADO
        celda.alignment = ALINEACION_ENCABEZADO
        celda.border = BORDE
        columna = (estados or {}).get(nombre)
        color = COLORES[columna.estado] if columna else "DDEBF7"
        celda.fill = PatternFill("solid", fgColor=color)
        if columna:
            nota = Comment(f"{columna.estado}\nOrigen: {columna.origen}\n{columna.regla}", "SAML")
            nota.width, nota.height = 320, 140
            celda.comment = nota
        hoja.column_dimensions[get_column_letter(indice)].width = ANCHOS.get(nombre, max(12, min(len(nombre) + 4, 24)))
    hoja.row_dimensions[1].height = 45

    for fila in filas:
        hoja.append([_celda(fila.get(nombre)) for nombre in columnas])

    for indice, nombre in enumerate(columnas, start=1):
        formato = FORMATOS.get(nombre)
        letra = get_column_letter(indice)
        for celda in hoja[letra][1:]:
            celda.alignment = Alignment(vertical="top")
            if formato:
                celda.number_format = formato
            if nombre == "Enlace" and isinstance(celda.value, str) and celda.value.startswith("https://"):
                celda.hyperlink = celda.value
                celda.style = "Hyperlink"

    hoja.freeze_panes = congelar
    hoja.auto_filter.ref = hoja.dimensions


def _hoja_resumen(
    hoja: Worksheet, resultado: lim.Resultado, generado: datetime, entrada: Path, modo: str | None = None
) -> None:
    m = resultado.metricas
    hoja.column_dimensions["A"].width = 44
    hoja.column_dimensions["B"].width = 34
    hoja.column_dimensions["C"].width = 24
    hoja.column_dimensions["D"].width = 30
    hoja.column_dimensions["E"].width = 80

    def titulo(texto: str) -> None:
        hoja.append([])
        hoja.append([texto])
        hoja.cell(row=hoja.max_row, column=1).font = FUENTE_SECCION

    def fila(etiqueta: str, valor: Any, formato: str | None = None) -> None:
        hoja.append([etiqueta, _celda(valor)])
        hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True)
        celda = hoja.cell(row=hoja.max_row, column=2)
        celda.alignment = Alignment(horizontal="left", vertical="top", wrap_text=isinstance(valor, str) and len(valor) > 40)
        if formato:
            celda.number_format = formato

    def encabezado(*columnas: str) -> None:
        hoja.append(list(columnas))
        for indice in range(1, len(columnas) + 1):
            celda = hoja.cell(row=hoja.max_row, column=indice)
            celda.font = FUENTE_ENCABEZADO
            celda.fill = PatternFill("solid", fgColor="DDEBF7")
            celda.border = BORDE

    hoja.append(["Excel intermedio de licitaciones adjudicadas"])
    hoja["A1"].font = FUENTE_TITULO
    hoja.append(
        [
            "Versión de demostración y validación. No es el producto final de GRECO y no reemplaza "
            "la integración con Obras V6, Supabase ni n8n."
        ]
    )
    hoja["A2"].font = Font(italic=True)

    titulo("Ejecución")
    fila("Generado el", generado, "dd/mm/yyyy hh:mm:ss")
    if modo:
        fila("Modo de ejecución", modo)
    entrada = entrada.resolve()
    fila("Carpeta de entrada", str(entrada.relative_to(RAIZ)) if entrada.is_relative_to(RAIZ) else str(entrada))
    recientes = [a.fecha_extraccion for a in resultado.archivos if a.fecha_extraccion]
    fila("Fecha de extracción (CSV más reciente)", max(recientes) if recientes else "No se pudo leer del nombre", "dd/mm/yyyy hh:mm:ss")
    fila("Archivos consolidados", len(resultado.archivos))
    fila(
        "Alcance de la consolidación",
        "Solo se consolidaron los CSV presentes en la carpeta de entrada al generar este archivo. "
        "Ante un idContrato repetido se usa la versión del CSV más reciente.",
    )
    encabezado("Archivo", "Fecha de extracción", "Filas", "SHA-256 (sin modificar)")
    for archivo in resultado.archivos:
        hoja.append([archivo.ruta.name, archivo.fecha_extraccion, archivo.filas, archivo.sha256])
        hoja.cell(row=hoja.max_row, column=2).number_format = "dd/mm/yyyy hh:mm:ss"

    titulo("Registros")
    fila("Filas leídas de los CSV", m["filas_leidas"])
    fila("Procesos únicos por idContrato", m["procesos_unicos"])
    fila("Duplicados dentro de un mismo CSV", m["duplicados_mismo_archivo"])
    fila("Reemplazados por una versión más reciente", m["reemplazados_por_version_reciente"])
    fila("Filas sin idContrato (excluidas)", m["filas_sin_id"])
    fila("Estados encontrados", ", ".join(f"{estado}: {n}" for estado, n in sorted(m["estados"].items())))
    fila("Filtro aplicado", f"estado_agrupado = '{lim.ESTADO_ADJUDICADO}'")
    todos_adjudicados = all(d["Estado (origen)"] == lim.ESTADO_ADJUDICADO for d in resultado.datos)
    fila(
        "Confirmación del filtro",
        f"Sí: los {m['incluidos']} registros de Datos tienen estado '{lim.ESTADO_ADJUDICADO}'."
        if todos_adjudicados
        else "NO: hay registros con otro estado; revisar.",
    )
    fila("Excluidos por estado", m["excluidos_por_estado"])
    fila("Registros incluidos en Datos", m["incluidos"])
    fila("idContrato únicos en Datos", "Sí" if m["ids_unicos_en_datos"] else "NO")
    fila("CodigoProceso repetidos", m["codigos_duplicados"])
    fila("Filas en Ganadores", m["filas_ganadores"])
    fila("Procesos con varios ganadores", m["procesos_varios_ganadores"])
    fila("Procesos sin ganador en origen", m["procesos_sin_ganador"])

    titulo("Unidad de Valor")
    fila("Valor mínimo distinto de 0 (pesos)", m["valor_minimo_no_cero"], "#,##0")
    fila("Mediana del Valor (pesos)", m["valor_mediana"], "#,##0")
    fila(
        "Conversión",
        "El campo Valor de la API está en pesos colombianos (si estuviera en millones, las cifras serían "
        "imposibles). 'Valor Millones' = Valor / 1.000.000; el valor original se conserva en "
        "'Valor (origen, pesos)'.",
    )

    titulo("Incidencias por tipo")
    encabezado("Tipo", "Cantidad")
    for tipo, cantidad in sorted(m["incidencias_por_tipo"].items(), key=lambda par: -par[1]):
        hoja.append([tipo, cantidad])
    fila("Total", len(resultado.incidencias))

    titulo("Campos manuales y calculados")
    manuales = [c.nombre for c in lim.COLUMNAS_GRECO if c.estado == lim.MANUAL]
    calculados = [c.nombre for c in lim.COLUMNAS_GRECO if c.estado == lim.CALCULADO]
    fila("Manuales (vacíos)", ", ".join(manuales))
    fila("Calculados (vacíos)", ", ".join(calculados))
    fila(
        "Motivo",
        "Los manuales los diligencia el equipo comercial y los calculados dependen de información o "
        "fórmulas del archivo original de GRECO (Obras V6), que todavía no se ha revisado.",
    )

    titulo("Campos del CSV que no se exportan")
    fila(
        "Excluidos",
        ", ".join(lim.CAMPOS_EXCLUIDOS)
        + ". Son identificadores firmados de la plataforma o estado del usuario en licitaciones.info; "
        "el CSV original los conserva.",
    )

    titulo("Leyenda de colores de encabezados")
    for estado, color in COLORES.items():
        hoja.append([estado, SIGNIFICADO_COLORES[estado]])
        hoja.cell(row=hoja.max_row, column=1).fill = PatternFill("solid", fgColor=color)

    titulo("Mapeo de columnas")
    encabezado("Hoja", "Columna", "Estado", "Origen", "Regla")
    for nombre_hoja, columnas in (("Datos", lim.COLUMNAS_DATOS), ("Ganadores", lim.COLUMNAS_GANADORES)):
        for columna in columnas:
            hoja.append([nombre_hoja, columna.nombre, columna.estado, columna.origen, columna.regla])
            hoja.cell(row=hoja.max_row, column=3).fill = PatternFill("solid", fgColor=COLORES[columna.estado])
            hoja.cell(row=hoja.max_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")


def construir_libro(
    resultado: lim.Resultado, generado: datetime, entrada: Path, modo: str | None = None
) -> Workbook:
    libro = Workbook()

    datos = libro.active
    datos.title = "Datos"
    _escribir_tabla(
        datos,
        [c.nombre for c in lim.COLUMNAS_DATOS],
        resultado.datos,
        estados={c.nombre: c for c in lim.COLUMNAS_DATOS},
        congelar="B2",
    )

    _escribir_tabla(
        libro.create_sheet("Ganadores"),
        [c.nombre for c in lim.COLUMNAS_GANADORES],
        resultado.ganadores,
        estados={c.nombre: c for c in lim.COLUMNAS_GANADORES},
    )

    incidencias = sorted(
        resultado.incidencias,
        key=lambda i: (i["Alcance"] != "Columna", i["Tipo"], i["idContrato"], i["Campo"]),
    )
    _escribir_tabla(libro.create_sheet("Incidencias"), list(lim.COLUMNAS_INCIDENCIAS), incidencias)

    _hoja_resumen(libro.create_sheet("Resumen"), resultado, generado, entrada, modo)
    return libro


def generar(entrada: Path, salida: Path, modo: str | None = None) -> tuple[Path, lim.Resultado]:
    """Limpia los CSV de `entrada` y guarda un Excel nuevo en `salida`.

    Lanza FileNotFoundError si no hay CSV y FileExistsError si el archivo de
    destino ya existe (nunca se sobrescribe).
    """
    resultado = lim.procesar(entrada)
    generado = datetime.now().replace(microsecond=0)
    salida.mkdir(parents=True, exist_ok=True)
    destino = salida / f"licitaciones_intermedio_{generado:%Y%m%d_%H%M%S}.xlsx"
    libro = construir_libro(resultado, generado, entrada, modo)
    with destino.open("xb") as archivo:  # "x": falla si ya existe
        libro.save(archivo)
    return destino, resultado


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--entrada", type=Path, default=ENTRADA_POR_DEFECTO, help="carpeta con los CSV de la ingesta")
    parser.add_argument("--salida", type=Path, default=SALIDA_POR_DEFECTO, help="carpeta donde se guarda el Excel")
    args = parser.parse_args()

    try:
        destino, resultado = generar(args.entrada, args.salida)
    except FileNotFoundError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    except FileExistsError as exc:
        print(f"ERROR: {exc.filename} ya existe; no se sobrescribe. Vuelve a ejecutar en unos segundos.", file=sys.stderr)
        return 1

    m = resultado.metricas
    print(f"CSV consolidados: {len(resultado.archivos)} | filas leidas: {m['filas_leidas']}")
    print(f"Licitaciones adjudicadas en Datos: {m['incluidos']} | excluidas por estado: {m['excluidos_por_estado']}")
    print(f"Filas en Ganadores: {m['filas_ganadores']} | incidencias: {len(resultado.incidencias)}")
    print(f"Excel generado: {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
